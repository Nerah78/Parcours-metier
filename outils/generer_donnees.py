#!/usr/bin/env python3
"""
Convertit le classeur cuisine_wow_forever.xlsx en Donnees.lua pour l'addon CuisineParcours.

Usage :
    python generer_donnees.py [classeur.xlsx] [Donnees.lua] [dossier_metiers]

La cuisine vient du classeur ; les autres métiers viennent des fichiers JSON du dossier « metiers »
(un fichier par métier : alchimie.json, forge.json, ...).
Par défaut : cuisine_wow_forever.xlsx + metiers/ (dossier courant) -> Parcours/Donnees.lua
Prérequis : pip install openpyxl
Important : le classeur doit contenir les valeurs calculées des formules (il suffit de l'ouvrir
puis de l'enregistrer une fois dans Excel ou LibreOffice).
Relancer le script après chaque modification du classeur.
"""
import glob
import json
import os
import re
import sys
from datetime import date
from openpyxl import load_workbook

CLASSEUR = sys.argv[1] if len(sys.argv) > 1 else "cuisine_wow_forever.xlsx"
SORTIE = sys.argv[2] if len(sys.argv) > 2 else "Parcours/Donnees.lua"
DOSSIER_METIERS = sys.argv[3] if len(sys.argv) > 3 else "metiers"

# Noms de zones (anglais -> français) pour l'affichage. Traductions à vérifier en jeu.
ZONES_FR = {
    "Tirisfal Glades": "Clairières de Tirisfal", "Undercity": "Fossoyeuse",
    "Darkshore": "Sombrivage", "Loch Modan": "Loch Modan", "Silverpine Forest": "Forêt des Pins argentés",
    "Ashenvale": "Orneval", "Hillsbrad Foothills": "Contreforts de Hautebrande",
    "Arathi Highlands": "Hautes-terres d'Arathi", "Desolace": "Désolace", "Badlands": "Terres ingrates",
    "Tanaris": "Tanaris", "Elwynn Forest": "Forêt d'Elwynn", "Durotar": "Durotar",
    "Westfall": "Marche de l'Ouest", "Dun Morogh": "Dun Morogh",
    "Redridge Mountains": "Les Carmines", "The Barrens": "Les Tarides",
    "Stranglethorn Vale": "Vallée de Strangleronce", "Winterspring": "Berceau-de-l'Hiver",
    "Duskwood": "Bois de la Pénombre", "Mulgore": "Mulgore", "Dustwallow Marsh": "Marécage d'Âprefange",
    "Teldrassil": "Teldrassil", "Feralas": "Féralas", "Swamp of Sorrows": "Marais des Chagrins",
    "Stonetalon Mountains": "Les Serres-Rocheuses", "Wetlands": "Les Paluns",
    "Un'Goro Crater": "Cratère d'Un'Goro", "The Hinterlands": "Les Hinterlands", "Azshara": "Azshara",
    "Felwood": "Gangrebois", "Western Plaguelands": "Maleterres de l'ouest",
    "Thousand Needles": "Mille pointes", "Alterac Mountains": "Montagnes d'Alterac",
    "Stormwind City": "Hurlevent", "All Faction Cities": "Toutes les capitales",
    "Capital cities": "Capitales", "All 1-10 zones": "Zones de niveau 1-10",
    "Thunder Bluff": "Les Pitons-du-Tonnerre", "Ironforge": "Forgefer", "Stormwind": "Hurlevent",
    "Booty Bay": "Baie-du-Butin", "Orgrimmar": "Orgrimmar", "Darnassus": "Darnassus",
}


# Zones QuestieDB mal traduites (noms de capitales anglais ou non traduits) -> nom français du client
CORRECTION_ZONES_QDB = {"Cité de Stormwind": "Hurlevent", "Thunder Bluff": "Les Pitons-du-Tonnerre",
                        "Ironforge": "Forgefer", "Undercity": "Fossoyeuse"}


def fr(zone):
    return ZONES_FR.get(zone, zone) if zone else None


def lire(ws, colonnes, ligne_min=2):
    """Retourne les lignes dont la 1re colonne est un entier, sous forme de dictionnaires."""
    lignes = []
    for row in ws.iter_rows(min_row=ligne_min, values_only=True):
        if isinstance(row[0], (int, float)) and not isinstance(row[0], bool):
            lignes.append({nom: row[i] if i < len(row) else None for i, nom in enumerate(colonnes)})
    return lignes


def entier(v):
    return int(v) if isinstance(v, float) and v == int(v) else v


# ---------------------------------------------------------------- lecture du classeur
wb = load_workbook(CLASSEUR, data_only=True)

recettes = {r["id"]: r for r in lire(wb["Recettes"], [
    "id", "type", "nom", "nouveau", "maitre", "qte", "orange", "jaune", "vert", "gris", "brut", "controle", "nbcomp"])}
plan = lire(wb["Plan_quêtes"], [
    "n", "rang", "debut", "fin", "recette", "recette_id", "couleur", "chance", "crafts", "obtention", "intitule", "sans_source", "obtention_detail"])
materiaux = lire(wb["Plan_matériaux"], [
    "n", "recette", "ing_id", "ingredient", "par_craft", "crafts", "total", "type_source", "zone", "statut"])
sources = lire(wb["Sources"], [
    "ing_id", "ingredient", "type", "source", "zone", "niveau_min", "niveau_max", "taux", "ref", "notes"])
marchands_lignes = lire(wb["Marchands"], ["npc", "nom", "zone", "map_id", "x", "y", "ing_id", "ingredient"])
maitres_qdb = lire(wb["QDB_Maitres"], ["npc", "nom", "titre", "titre_en", "zone", "coord", "faction", "niveau"]) \
    if "QDB_Maitres" in wb.sheetnames else []
recettes_vendues = lire(wb["Marchands_recettes"], [
    "npc", "nom", "zone", "map_id", "x", "y", "objet_id", "objet", "recette_id", "recette", "origine"])

if not plan:
    sys.exit("Plan_quêtes est vide : ouvre puis enregistre le classeur pour calculer les formules.")
if any(p["crafts"] is None for p in plan):
    sys.exit("Valeurs calculées absentes : ouvre le classeur dans Excel/LibreOffice et enregistre-le.")

# ---------------------------------------------------------------- construction des données
ing_du_plan = {m["ing_id"] for m in materiaux}

# recette cuisinée -> objet "Recette : ..." vendu et marchand (relevés en jeu)
vendue_par = {}
for l in recettes_vendues:
    vendue_par.setdefault(l["recette_id"], l)

quetes = []
for p in plan:
    rec = recettes.get(p["recette_id"], {})
    vente = vendue_par.get(p["recette_id"])
    ingredients = [
        {"id": m["ing_id"], "nom": m["ingredient"], "par_craft": entier(m["par_craft"])}
        for m in materiaux if m["n"] == p["n"]
    ]
    quetes.append({
        "n": entier(p["n"]), "rang": p["rang"], "debut": entier(p["debut"]), "fin": entier(p["fin"]),
        "couleur": p["couleur"], "chance": p["chance"], "crafts": entier(p["crafts"]),
        "obtention": p["obtention"], "obtentionDetail": p.get("obtention_detail"),
        "recette": {
            "id": p["recette_id"], "nom": p["recette"], "maitre": rec.get("maitre") == "Oui",
            "orange": rec.get("orange"), "jaune": rec.get("jaune"), "vert": rec.get("vert"), "gris": rec.get("gris"),
            "objetId": entier(vente["objet_id"]) if vente else None,
            "vendeur": ({"nom": vente["nom"], "zone": vente["zone"], "mapID": entier(vente["map_id"]),
                         "x": vente["x"], "y": vente["y"],
                         "origine": "web" if str(vente.get("origine") or "").startswith("Web") else None} if vente else None),
        },
        "ingredients": ingredients,
    })

ingredients = {}
for ing_id in sorted(ing_du_plan):
    nom = next(m["ingredient"] for m in materiaux if m["ing_id"] == ing_id)
    ingredients[ing_id] = {"nom": nom, "sources": []}
for s in sources:
    if s["ing_id"] in ingredients:
        ingredients[s["ing_id"]]["sources"].append({
            "type": s["type"], "source": s["source"], "zone": s["zone"], "zoneFr": fr(s["zone"]),
            "niveauMin": entier(s["niveau_min"]), "niveauMax": entier(s["niveau_max"]),
            "taux": entier(s["taux"]), "notes": s["notes"],
        })

marchands = {}
for l in marchands_lignes:
    m = marchands.setdefault(l["npc"], {
        "npcId": entier(l["npc"]), "nom": l["nom"], "zone": l["zone"], "mapID": entier(l["map_id"]),
        "x": l["x"], "y": l["y"], "composants": [], "recettes": []})
    if l["ing_id"] not in m["composants"]:
        m["composants"].append(entier(l["ing_id"]))
for l in recettes_vendues:
    m = marchands.setdefault(l["npc"], {
        "npcId": entier(l["npc"]), "nom": l["nom"], "zone": l["zone"], "mapID": entier(l["map_id"]),
        "x": l["x"], "y": l["y"], "composants": [], "recettes": []})
    m["recettes"].append({"objetId": entier(l["objet_id"]), "objet": l["objet"],
                          "recetteId": entier(l["recette_id"]), "recette": l["recette"]})

maitres = []
for m in maitres_qdb:
    try:
        x, y = [float(t) for t in str(m["coord"]).split(",")]
    except ValueError:
        x = y = None
    zone = CORRECTION_ZONES_QDB.get(m["zone"], m["zone"])
    maitres.append({"npcId": entier(m["npc"]), "nom": m["nom"], "titre": m["titre"], "zone": zone,
                    "x": x, "y": y, "faction": m["faction"], "niveau": entier(m["niveau"]), "origine": "questie"})

cuisine = {
    "id": "cuisine", "nom": "Cuisine", "motsClient": ["cuisine", "cooking"], "ligneId": 185,
    "icone": "INV_Misc_Food_15", "edition": "Forever (plan du tableur)",
    "quetes": quetes, "ingredients": ingredients, "maitres": maitres,
    "marchands": sorted(marchands.values(), key=lambda m: (m["zone"] or "", m["nom"])),
}

# ---------------------------------------------------------------- autres métiers (fichiers JSON)
# Identité de chaque métier : noms possibles dans le client (minuscules), id de compétence, icône
METIERS = {
    "alchimie": ("Alchimie", ["alchimie", "alchemy"], 171, "Trade_Alchemy"),
    "forge": ("Forge", ["forge", "blacksmith"], 164, "Trade_BlackSmithing"),
    "ingenierie": ("Ingénierie", ["ingénierie", "ingenierie", "engineering"], 202, "Trade_Engineering"),
    "couture": ("Couture", ["couture", "tailoring"], 197, "Trade_Tailoring"),
    "cuir": ("Travail du cuir", ["travail du cuir", "leatherworking"], 165, "Trade_LeatherWorking"),
    "enchantement": ("Enchantement", ["enchantement", "enchanting"], 333, "Trade_Engraving"),
    "herboristerie": ("Herboristerie", ["herboristerie", "herbalism"], 182, "Trade_Herbalism"),
    "minage": ("Minage", ["minage", "mining"], 186, "Trade_Mining"),
    "depecage": ("Dépeçage", ["dépeçage", "depecage", "skinning"], 393, "INV_Misc_Pelt_Wolf_01"),
    "secourisme": ("Secourisme", ["secourisme", "first aid"], 129, "Spell_Holy_SealOfSacrifice"),
    "peche": ("Pêche", ["pêche", "peche", "fishing"], 356, "Trade_Fishing"),
}
ORDRE = ["cuisine", "alchimie", "forge", "ingenierie", "couture", "cuir", "enchantement",
         "herboristerie", "minage", "depecage", "secourisme", "peche"]
# Métiers de récolte (pas de recette) : marqués « recolte », masqués par défaut dans l'addon (option dans le journal)
EXCLUS = {"herboristerie", "minage", "depecage", "peche"}


def etape_vers_quete(e):
    """Étape d'un plan JSON -> quête au même format que la cuisine."""
    crafts = e.get("crafts")
    variantes = None
    if isinstance(crafts, list):  # plusieurs sorts dans la même étape (fonte, bandages)
        variantes = [{"nom": c.get("nomEn"), "sortId": c.get("sortId"), "quantite": c.get("quantite")} for c in crafts]
        crafts = None
    obtention = e.get("obtention") or "Inconnue"
    recette = None
    if e.get("recetteEn") or e.get("recetteSortId"):
        recette = {"id": e.get("recetteSortId"), "sortId": e.get("recetteSortId"), "objetId": e.get("recetteObjetId"),
                   "nom": e.get("recetteEn"), "maitre": obtention == "Maître"}
        if obtention.startswith("Marchand"):
            nom_marchand, _, zone = obtention.partition(":")[2].strip().partition(",")
            recette["vendeur"] = {"nom": nom_marchand.strip(), "zone": zone.strip() or None, "origine": "web"}
    nom_variantes = ", ".join(v["nom"] for v in variantes if v["nom"]) if variantes else None
    ings = []
    for i in e.get("ingredients") or []:
        par = i.get("parCraft")
        if par is None and i.get("total") and crafts:
            par = round(i["total"] / crafts, 2)
        ings.append({"id": i.get("id"), "nom": i.get("nomEn"), "par_craft": entier(par)})
    titre = (recette or {}).get("nom") or nom_variantes
    return {
        "n": e["n"], "debut": e["debut"], "fin": e["fin"], "crafts": entier(crafts),
        "obtention": obtention, "obtentionDetail": e.get("notes"), "titre": titre,
        "recette": recette, "variantes": variantes, "zones": e.get("zones"), "ingredients": ings,
    }


metiers = {"cuisine": cuisine}
for chemin in sorted(glob.glob(os.path.join(DOSSIER_METIERS, "*.json"))):
    j = json.load(open(chemin, encoding="utf-8"))
    mid = j["id"]
    if mid not in METIERS:
        print(f"Métier ignoré (inconnu) : {mid}")
        continue
    nom, mots, ligne, icone = METIERS[mid]
    quetes_m = [etape_vers_quete(e) for e in j["etapes"]]
    for q in quetes_m:
        if not q["titre"]:
            q["titre"] = f"{nom} {q['debut']} - {q['fin']}"
    maitres_m = [{"nom": m["nom"], "zone": fr(m.get("zone")), "faction": m.get("faction"), "x": m.get("x"), "y": m.get("y"),
                  "origine": "web"} for m in j.get("maitres", [])]
    metiers[mid] = {
        "id": mid, "nom": nom, "recolte": (mid in EXCLUS) or None, "motsClient": mots, "ligneId": ligne, "icone": icone,
        "edition": j.get("edition"), "source": j.get("source"), "limites": j.get("limites"),
        "quetes": quetes_m, "ingredients": {}, "maitres": maitres_m, "marchands": [],
    }

donnees = {
    "meta": {"genere": date.today().isoformat(), "classeur": CLASSEUR, "nbMetiers": len(metiers)},
    "ordre": [m for m in ORDRE if m in metiers],
    "metiers": metiers,
}


# ---------------------------------------------------------------- sérialisation Lua
def cle(k):
    if isinstance(k, int):
        return f"[{k}]"
    return k if re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", k) else f'["{k}"]'


def chaine(s):
    s = s.replace("\u00a0", " ")  # espaces insécables (présents dans les noms QuestieDB)
    return '"' + s.replace("\\", "\\\\").replace('"', '\\"').replace("\n", "\\n").replace("\r", "") + '"'


def lua(v, niveau=0):
    pad, pad2 = "  " * niveau, "  " * (niveau + 1)
    if v is None:
        return "nil"
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        return str(int(v)) if v == int(v) else repr(round(v, 4))
    if isinstance(v, str):
        return chaine(v)
    if isinstance(v, (list, tuple)):
        if not v:
            return "{}"
        return "{\n" + "".join(f"{pad2}{lua(x, niveau + 1)},\n" for x in v) + pad + "}"
    if isinstance(v, dict):
        items = [(k, x) for k, x in v.items() if x is not None]
        if not items:
            return "{}"
        return "{\n" + "".join(f"{pad2}{cle(k)} = {lua(x, niveau + 1)},\n" for k, x in items) + pad + "}"
    raise TypeError(type(v))


with open(SORTIE, "w", encoding="utf-8") as f:
    f.write("-- Fichier généré par generer_donnees.py : NE PAS MODIFIER À LA MAIN\n")
    f.write(f"-- Source : {CLASSEUR} - généré le {date.today().isoformat()}\n")
    f.write("ParcoursData = " + lua(donnees) + "\n")

print(f"{SORTIE} : {len(metiers)} métiers (" + ", ".join(f"{k}:{len(v['quetes'])}" for k, v in metiers.items()) + ")")
